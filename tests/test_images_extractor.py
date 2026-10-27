from __future__ import annotations

from pathlib import Path

import pytest
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

from medha.images import DoclingImageExtractor
from medha.ingest.docling_parser import missing_artifacts


def _make_pptx(path: Path) -> None:
    picture = path.parent / "diagram.png"
    image = Image.new("RGB", (640, 360), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle([40, 40, 600, 320], outline="black", width=4)
    draw.line([60, 300, 280, 100], fill="red", width=6)
    draw.text((320, 170), "F = ma", fill="black")
    image.save(picture)

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(str(picture), Inches(1), Inches(1), width=Inches(6))
    presentation.save(str(path))


def test_extracts_pptx_picture(tmp_path: Path) -> None:
    pytest.importorskip("docling")
    if missing_artifacts():
        pytest.skip("Docling artifacts are not prefetched")
    source = tmp_path / "lesson.pptx"
    _make_pptx(source)
    output = tmp_path / "images"
    result = DoclingImageExtractor(output_dir=output, num_threads=2).extract(source)
    assert result.images
    image = result.images[0]
    assert image.file.endswith(".png")
    assert (output / image.file).exists()
    assert image.width >= 64
    assert image.height >= 64
    assert image.page_no == 1
    assert image.locator_kind == "slide"
