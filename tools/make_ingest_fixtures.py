"""Generate synthetic PDF and PPTX fixtures for ingestion tests."""

from __future__ import annotations

import argparse
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUT = ROOT / "tests" / "fixtures" / "ingest"


def _pdf(out: Path) -> Path:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages

    out.parent.mkdir(parents=True, exist_ok=True)
    with PdfPages(out) as pdf:
        fig = plt.figure(figsize=(8.5, 11))
        fig.patch.set_facecolor("white")
        fig.text(0.09, 0.92, "Newton's Laws of Motion", fontsize=24, weight="bold")
        fig.text(0.09, 0.86, "Physics 101", fontsize=13, color="#444444")
        fig.text(0.09, 0.78, "Newton's Second Law", fontsize=17, weight="bold")
        body = (
            "The acceleration of an object depends on the net force acting on it "
            "and the mass of the object.\n\n"
            "A larger net force produces a larger acceleration. For the same force, "
            "a larger mass produces a smaller acceleration."
        )
        fig.text(0.09, 0.72, body, fontsize=11, va="top", linespacing=1.5)
        fig.text(0.09, 0.50, "F = m a", fontsize=19, family="DejaVu Sans")
        fig.text(
            0.09,
            0.43,
            "Force is measured in newtons, mass in kilograms, and acceleration "
            "in metres per second squared.",
            fontsize=11,
            va="top",
            linespacing=1.5,
        )
        pdf.savefig(fig, bbox_inches="tight")
        plt.close(fig)

        fig = plt.figure(figsize=(8.5, 11))
        fig.patch.set_facecolor("white")
        fig.text(0.09, 0.92, "Worked Example", fontsize=20, weight="bold")
        fig.text(0.09, 0.84, "A 2 kg block is pushed by a net force of 10 N.", fontsize=11)
        fig.text(0.09, 0.78, "Step 1: Write the equation.", fontsize=11, weight="bold")
        fig.text(0.09, 0.73, "a = F / m", fontsize=17)
        fig.text(0.09, 0.65, "Step 2: Substitute the values.", fontsize=11, weight="bold")
        fig.text(0.09, 0.60, "a = 10 N / 2 kg = 5 m/s^2", fontsize=17)
        fig.text(0.09, 0.50, "The block accelerates at 5 metres per second squared.", fontsize=11)
        pdf.savefig(fig, bbox_inches="tight")
        plt.close(fig)
    return out


def _add_bullets(text_frame: Any, items: list[tuple[str, int]]) -> None:
    frame = text_frame
    frame.clear()
    for index, (text, level) in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = text
        paragraph.level = level


def _pptx(out: Path) -> Path:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    out.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Newton's Second Law"
    slide.placeholders[1].text = "Force, mass, and acceleration"

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Key quantities"
    _add_bullets(
        slide.placeholders[1].text_frame,
        [
            ("Force is measured in newtons (N).", 0),
            ("Mass is measured in kilograms (kg).", 0),
            ("Acceleration is measured in m/s^2.", 0),
        ],
    )
    rows, cols = 3, 3
    table_shape = slide.shapes.add_table(
        rows, cols, Inches(7.0), Inches(2.0), Inches(5.3), Inches(1.6)
    )
    table = table_shape.table
    data = [
        ("Quantity", "Symbol", "Unit"),
        ("Force", "F", "newton"),
        ("Mass", "m", "kilogram"),
    ]
    for row_index, row in enumerate(data):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.text = value
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(16)

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Worked example"
    _add_bullets(
        slide.placeholders[1].text_frame,
        [
            ("A 2 kg block is pushed by a net force of 10 N.", 0),
            ("Use a = F / m.", 0),
            ("Substitute: a = 10 / 2 = 5 m/s^2.", 1),
        ],
    )
    slide.notes_slide.notes_text_frame.text = "Speaker note: this line must not be ingested."

    prs.save(str(out))
    return out


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(prog="make-ingest-fixtures", description=__doc__)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = parser.parse_args(argv)

    out = args.out.resolve()
    print(_pdf(out / "sample.pdf"))
    print(_pptx(out / "sample.pptx"))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
